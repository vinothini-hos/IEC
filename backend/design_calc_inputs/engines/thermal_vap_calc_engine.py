"""
thermal_vap_calc_engine.py - thermal vap calculation. Self-contained: evaluates json/sample thermal vap.json
with the cl2 enquiry sheet.

Run:  python thermal_vap_calc_engine.py


Formula language supported (value strings starting with "="):
    =f_multiply(a, b, ...)   =f_div(a, b)   =f_add(a, b, ...)   =f_sub(a, b)
    =f_pow(a, b)  =f_min(...)  =f_max(...)  =f_round(a, n)  =f_roundof(a, n) round off like Excel ROUND (n = 0 if left out)
    =f_vlookup(value, $sheet1.VLOOKUP_001, col, 0)   Excel VLOOKUP on a lookup table JSON
    =f_vlookup(value, VLOOKUP1.B17.nb:VLOOKUP1.B20.od, col)   same, Excel-style range
    =VLOOKUP(...)            same as f_vlookup
                     (sheet_id.Vlookup_table_id; col = number or column name; 0 exact, 1 approx)
    =f_addition(...) = f_add   =f_division(a, b) = f_div
    =LN(a)  =LOG10(a)  =EXP(a)  =SQRT(a)  =ABS(a)
    =f_def1(...), =f_def2(...), ...  formulas from the def JSON ("sample thermal vap def.json"),
                     arguments a, b, c, ... in order - see load_defs
    plain arithmetic: + - * / ^ ( )  e.g. =@T.D7*2.205
    references:
        @TABLE_ID.D7        -> cell D7 in table TABLE_ID (column letter -> value key via COLUMN_MAP)
        @TABLE_ID.value1.B7 -> explicit key + row id (alternative form)
        @TABLE_ID.B7.value1 -> same, row id first
        sheet_name.field    -> external input, e.g. enquiry_sheet.abc (supplied via `inputs`)
    literals: numbers, "quoted strings"
Anything not starting with "=" is a literal ("-" or "" = empty, "???" = missing input).
"""
import json, csv, re, math
from pathlib import Path

# Excel column letter -> JSON value key. Row id "B7" -> row number 7.
COLUMN_MAP = {"D": "value1", "E": "value2"}
VALUE_KEYS = ["value1", "value2", "value"]
KEY_COL = {"value1": "D", "value2": "E", "value": "D"}   # "value": single-value tables
HEADER_COL = {"description": "B", "unit1": "C", "unit": "C", "value1": "D", "value": "D",
              "value2": "E", "unit2": "F", "notes": "G"}      # table column key -> Excel column

def _math_fn(name, fn):
    """A math function that reports a bad argument (LN(0), SQRT(-1), text) as a CalcError."""
    def call(*args):
        for v in args:
            if not isinstance(v, (int, float)):
                raise CalcError(f"{name}: argument is text ({v!r}), expected a number")
        try:
            return fn(*args)
        except (ValueError, OverflowError):
            raise CalcError(f"{name}({', '.join(f'{v:.6g}' for v in args)}) is undefined")
    return call

def f_power(a, b):
    """a ^ b, as Excel: a negative base with a fractional exponent is an error, not a complex number."""
    for v in (a, b):
        if not isinstance(v, (int, float)):
            raise CalcError(f"^: argument is text ({v!r}), expected a number")
    if a < 0 and not float(b).is_integer():
        raise CalcError(f"{a:.6g}^{b:.6g} is undefined")
    return a ** b

# ---------------- def functions (f_def1, f_def2, ...) ----------------
# Loaded from the sheet's def JSON ("sample thermal vap def.json"):
#     {"functions": {"f_def1": {"formula": "a*2.205"}, "f_def3": {"formula": "(a+b)/2"}, ...}}
# The arguments of a call are a, b, c, ... in order, so =f_def3(D10, D11) is (D10+D11)/2.
# A def can instead list its own names: {"params": ["od", "thk"], "formula": "(od-(2*thk))/2"}.
DEFS = {}         # "f_def1" -> {"params": [...], "formula": "...", "tokens": [...]}

def load_defs(path):
    """Load the "functions" of a def JSON into DEFS, checking each formula parses and uses only its params."""
    DEFS.clear()
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    for name, spec in (data.get("functions") or {}).items():
        if not isinstance(spec, dict) or "formula" not in spec:
            raise CalcError(f"{Path(path).name}: {name} needs a \"formula\"")
        tokens = tokenize(str(spec["formula"]).lstrip("="))
        names = {v for i, (k, v) in enumerate(tokens)
                 if k == "name" and (i + 1 == len(tokens) or tokens[i + 1][1] != "(")}
        params = spec.get("params")
        if params is None:
            bad = sorted(n for n in names if not re.fullmatch(r"[a-z]", n))
            if bad:
                raise CalcError(f"{Path(path).name}: {name} uses {bad} - write arguments as a, b, c, ... "
                                f"or list them in \"params\"")
            params = [chr(c) for c in range(ord("a"), max(map(ord, names), default=ord("a") - 1) + 1)]
        unknown = sorted(names - set(params))
        if unknown:
            raise CalcError(f"{Path(path).name}: {name} uses {unknown}, not in its params {params}")
        DEFS[name] = {"params": list(params), "formula": spec["formula"], "tokens": tokens}
    return DEFS

def _def_env(name, args):
    d = DEFS[name]
    if len(args) != len(d["params"]):
        raise CalcError(f"{name} expects {len(d['params'])} argument(s) ({', '.join(d['params'])}), got {len(args)}")
    return d["tokens"], dict(zip(d["params"], args))

# ---------------- VLOOKUP tables ----------------
LOOKUPS = {}      # "$sheet1.VLOOKUP_001" -> lookup table JSON (sheet_id + Vlookup_table_id)
LOOKUP_XL = {}    # same key -> Excel range of that table, filled while writing the .xlsx

def load_lookups(folder):
    """Load every lookup table JSON (one with "Vlookup_table_id") found in `folder` into LOOKUPS."""
    for path in sorted(Path(folder).glob("*.json")):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (ValueError, UnicodeDecodeError):
            continue
        for t in data if isinstance(data, list) else [data]:
            if isinstance(t, dict) and "Vlookup_table_id" in t:
                key = f"{t.get('sheet_id')}.{t['Vlookup_table_id']}"
                if key in LOOKUPS:
                    print(f"warning: lookup {key} defined twice, using {path.name}")
                LOOKUPS[key] = t
    return LOOKUPS

def lookup_table(name):
    t = LOOKUPS.get(name)
    if t is None:
        raise CalcError(f"Lookup table {name} not loaded (loaded: {list(LOOKUPS)})")
    return t

def lookup_col(t, col):
    """Column number (1 = first column) from a number or a column key / display name."""
    cols = t.get("columns", [])
    if isinstance(col, (int, float)):
        i = int(col)
    else:
        i = next((n for n, c in enumerate(cols, 1) if col in (c["key"], c.get("display_name"))), 0)
    if not 1 <= i <= len(cols):
        raise CalcError(f"Column {col!r} not in lookup {t['Vlookup_table_id']} (1..{len(cols)})")
    return i

def lookup_cell(t, row, i):
    """Value of column i in a lookup row - stored under the column's key or its display_name."""
    c = t["columns"][i - 1]
    return row.get(c["key"], row.get(c.get("display_name")))

def lookup_range(spec):
    """Resolve a lookup reference to (key, table, rows, first column no., last column no.).
    spec: "$sheet1.VLOOKUP_001"                  whole table, by sheet id + lookup table id
          "VLOOKUP1.B17.nb:VLOOKUP1.B20.od"      Excel-style range: table name or id, rows, columns"""
    if ":" not in spec:
        t = lookup_table(spec)
        return spec, t, t["rows"], 1, len(t.get("columns", []))
    (n1, r1, c1), (n2, r2, c2) = (p.split(".") for p in spec.split(":"))
    if n1 != n2:
        raise CalcError(f"Range {spec} spans two lookup tables")
    keys = [k for k, t in LOOKUPS.items() if n1 in (t.get("Vlookup_table_name"), t["Vlookup_table_id"])]
    if not keys:
        raise CalcError(f"Lookup table {n1} not loaded (loaded: {list(LOOKUPS)})")
    if len(keys) > 1:
        raise CalcError(f"Lookup name {n1} is in several sheets {keys} - use $sheet_id.table_id")
    t = LOOKUPS[keys[0]]
    lo, hi = sorted((int(r1[1:]), int(r2[1:])))
    rows = [r for r in t["rows"] if lo <= int(r["id"][1:]) <= hi]
    ci, cj = sorted((lookup_col(t, c1), lookup_col(t, c2)))
    return keys[0], t, rows, ci, cj

def range_col(t, ci, cj, col, spec):
    """Absolute column no. of `col` (number counted from the range's first column, or a name)."""
    i = ci + int(col) - 1 if isinstance(col, (int, float)) else lookup_col(t, col)
    if not ci <= i <= cj:
        raise CalcError(f"Column {col!r} is outside {spec}")
    return i

def f_vlookup(value, table, col, approx=0):
    """Excel VLOOKUP(value, table, col, approx) on a lookup table JSON.
    table: $sheet1.VLOOKUP_001 (sheet id + lookup table id) or VLOOKUP1.B17.nb:VLOOKUP1.B20.od (range)
    col: column number counted from the first column of the table/range, or a column name
    approx: 0 = exact match (default), 1 = largest first-column value <= value (Excel TRUE)"""
    _, t, rows, ci, cj = lookup_range(table)
    i = range_col(t, ci, cj, col, table)
    match = None
    for r in rows:
        k = lookup_cell(t, r, ci)
        if not approx and k == value:
            match = r
            break
        if approx and isinstance(k, (int, float)) and k <= value and (
                match is None or k >= lookup_cell(t, match, ci)):
            match = r
    if match is None:
        raise CalcError(f"{value!r} not found in first column of {table}")
    v = lookup_cell(t, match, i)
    if v in (None, "-", ""):
        raise CalcError(f"{table}: no value in column {i} for {value!r}")
    return v

def xl_vlookup(value, table, col, approx="0"):
    """Excel form of f_vlookup: VLOOKUP(value, 'VLOOKUP1'!$B$17:$C$20, col, FALSE)."""
    from openpyxl.utils import get_column_letter as L
    key, t, rows, ci, cj = lookup_range(table)
    if key not in LOOKUP_XL:
        raise CalcError(f"Lookup table {table} not in the workbook")
    sheet, shift = LOOKUP_XL[key]
    nums = [int(r["id"][1:]) + shift for r in rows]
    rng = f"'{sheet}'!${L(ci + 1)}${min(nums)}:${L(cj + 1)}${max(nums)}"
    if col.startswith('"'):
        col = range_col(t, ci, cj, col[1:-1], table) - ci + 1
    elif re.fullmatch(r"\d+\.?\d*", col):
        col = range_col(t, ci, cj, float(col), table) - ci + 1
    exact = approx in ("0", "0.0", "FALSE")
    return f"VLOOKUP({value},{rng},{col},{'FALSE' if exact else 'TRUE'})"

FUNCS = {
    "f_multiply": lambda *a: math.prod(a),
    "f_div": lambda a, b: a / b,
    "f_add": lambda *a: sum(a),
    "f_sub": lambda a, b: a - b,
    "f_pow": lambda a, b: a ** b,
    "f_min": min,
    "f_max": max,
    "f_round": lambda a, n=0: round(a, int(n)),
    "f_roundof": lambda a, n=0: math.copysign(math.floor(abs(a) * 10 ** int(n) + 0.5), a) / 10 ** int(n),  # Excel ROUND: halves away from 0
    "f_vlookup": f_vlookup,          # VLOOKUP on a lookup table JSON
    "VLOOKUP": f_vlookup,            # same, Excel spelling: =VLOOKUP(value, VLOOKUP1.B36.nb:VLOOKUP1.B63.od, 2)
    # Excel math functions, for def formulas ("LN((a-b)/(c-d))") and calc JSONs alike
    "LN": _math_fn("LN", math.log),
    "LOG10": _math_fn("LOG10", math.log10),
    "EXP": _math_fn("EXP", math.exp),
    "SQRT": _math_fn("SQRT", math.sqrt),
    "ABS": _math_fn("ABS", abs),
    # f_def1, f_def2, ... come from the sheet's def JSON (DEFS), see load_defs
    "f_addition": lambda *a: sum(a),  # same as f_add
    "f_division": lambda a, b: a / b, # same as f_div
}

# Excel versions of FUNCS, used when writing formulas to the .xlsx (arguments arrive already bracketed)
EXCEL_FUNCS = {
    "f_multiply": lambda *a: "*".join(a),
    "f_div": lambda a, b: f"{a}/{b}",
    "f_add": lambda *a: "+".join(a),
    "f_sub": lambda a, b: f"{a}-{b}",
    "f_pow": lambda a, b: f"POWER({a},{b})",
    "f_min": lambda *a: f"MIN({','.join(a)})",
    "f_max": lambda *a: f"MAX({','.join(a)})",
    "f_round": lambda a, n="0": f"ROUND({a},{n})",
    "f_roundof": lambda a, n="0": f"ROUND({a},{n})",
    "f_vlookup": xl_vlookup,
    "VLOOKUP": xl_vlookup,
    "LN": lambda a: f"LN({xl_unwrap(a)})",          # the call's own brackets are enough
    "LOG10": lambda a: f"LOG10({xl_unwrap(a)})",
    "EXP": lambda a: f"EXP({xl_unwrap(a)})",
    "SQRT": lambda a: f"SQRT({xl_unwrap(a)})",
    "ABS": lambda a: f"ABS({xl_unwrap(a)})",
    "f_addition": lambda *a: "+".join(a),
    "f_division": lambda a, b: f"{a}/{b}",
}

class CalcError(Exception):
    pass

# ---------------- tokenizer + recursive-descent parser ----------------
TOKEN_RE = re.compile(r"""
    (?P<num>\d+\.?\d*(?:[eE][-+]?\d+)?|\.\d+)
  | (?P<str>"[^"]*")
  | (?P<rng>\w+\.\w+\.\w+:\w+\.\w+\.\w+)    # VLOOKUP1.B17.nb:VLOOKUP1.B20.od (lookup range)
  | (?P<lkp>\$\w+\.\w+)                         # $sheet1.VLOOKUP_001 (lookup table)
  | (?P<ref>@?[A-Za-z_][\w]*(?:\.[\w]+)+)      # @T.D7, @T.value1.B7, enquiry_sheet.abc
  | (?P<name>[A-Za-z_]\w*)                      # function names / bare cell refs
  | (?P<op>[-+*/^(),])
  | (?P<ws>\s+)
""", re.X)

def tokenize(s):
    pos, out = 0, []
    while pos < len(s):
        m = TOKEN_RE.match(s, pos)
        if not m:
            raise CalcError(f"Unexpected character {s[pos]!r} in formula {s!r}")
        pos = m.end()
        if m.lastgroup != "ws":
            out.append((m.lastgroup, m.group()))
    return out

class Parser:
    def __init__(self, tokens, resolve):
        self.t, self.i, self.resolve = tokens, 0, resolve
    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else (None, None)
    def take(self, val=None):
        tok = self.peek()
        if val is not None and tok[1] != val:
            raise CalcError(f"Expected {val!r}, got {tok[1]!r}")
        self.i += 1
        return tok
    def parse(self):
        v = self.expr()
        if self.i != len(self.t):
            raise CalcError(f"Trailing tokens: {self.t[self.i:]}")
        return v
    def expr(self):
        v = self.term()
        while self.peek()[1] in ("+", "-"):
            op = self.take()[1]; r = self.term()
            v = v + r if op == "+" else v - r
        return v
    def term(self):
        v = self.power()
        while self.peek()[1] in ("*", "/"):
            op = self.take()[1]; r = self.power()
            v = v * r if op == "*" else v / r
        return v
    def power(self):
        """a^b binds tighter than * and /, left to right like Excel (2^3^2 = 64; -2^2 = 4)."""
        v = self.factor()
        while self.peek()[1] == "^":
            self.take(); v = self.pow_op(v, self.factor())
        return v
    def pow_op(self, a, b):
        return f_power(a, b)
    def args(self):
        self.take("(")
        args = []
        if self.peek()[1] != ")":
            args.append(self.expr())
            while self.peek()[1] == ",":
                self.take(); args.append(self.expr())
        self.take(")")
        return args
    def call_def(self, name, args):
        """Evaluate a def formula with its params bound to the call's argument values."""
        tokens, env = _def_env(name, args)
        def resolve(n):
            if n not in env:
                raise CalcError(f"{name}: unknown name {n}")
            return env[n]
        return Parser(tokens, resolve).parse()
    def factor(self):
        kind, val = self.peek()
        if val == "-":
            self.take(); return -self.factor()
        if val == "+":
            self.take(); return self.factor()
        if val == "(":
            self.take(); v = self.expr(); self.take(")"); return v
        if kind == "num":
            self.take(); return float(val)
        if kind == "str":
            self.take(); return val[1:-1]
        if kind in ("lkp", "rng"):
            self.take(); return val
        if kind == "ref":
            self.take(); return self.resolve(val)
        if kind == "name":
            self.take()
            if self.peek()[1] == "(":            # function call
                if val in DEFS:
                    return self.call_def(val, self.args())
                if val not in FUNCS:
                    hint = " - is it in the sheet's def JSON?" if val.startswith("f_def") else ""
                    raise CalcError(f"Unknown function {val}{hint}")
                return FUNCS[val](*self.args())
            return self.resolve(val)             # bare cell ref like D7 (same table)
        raise CalcError(f"Unexpected token {val!r}")

XL_ATOM_RE = re.compile(r"(?:'[^']*'!|\w+!)?\$?[A-Z]+\$?\d+|\d+\.?\d*(?:[eE][-+]?\d+)?|\.\d+|\"[^\"]*\"|\$\w+\.\w+|\w+\.\w+\.\w+:\w+\.\w+\.\w+")

def xl_wrap(x):
    """Bracket an Excel sub-expression unless it is a single cell, number, string or function call."""
    if XL_ATOM_RE.fullmatch(x):
        return x
    m = re.match(r"[A-Z][A-Z0-9.]*\(", x)
    if m:
        depth = 0
        for i in range(m.end() - 1, len(x)):
            depth += (x[i] == "(") - (x[i] == ")")
            if depth == 0:
                return x if i == len(x) - 1 else f"({x})"
    return f"({x})"

def xl_unwrap(x):
    """Drop one pair of brackets around the whole formula: (D8*2.205) -> D8*2.205."""
    if x.startswith("(") and x.endswith(")"):
        depth = 0
        for i, ch in enumerate(x):
            depth += (ch == "(") - (ch == ")")
            if depth == 0:
                return x[1:-1] if i == len(x) - 1 else x
    return x

class ExcelFormula(Parser):
    """Same grammar as Parser, but builds an Excel formula string instead of a number.
    References are turned into cell addresses by `resolve`; f_ functions via EXCEL_FUNCS."""
    def expr(self):
        v = self.term()
        while self.peek()[1] in ("+", "-"):
            op = self.take()[1]; v = f"{v}{op}{self.term()}"
        return v
    def term(self):
        v = self.power()
        while self.peek()[1] in ("*", "/"):
            op = self.take()[1]; v = f"{v}{op}{self.power()}"
        return v
    def pow_op(self, a, b):
        return f"{a}^{b}"
    def call_def(self, name, args):
        """The def formula as Excel, with its params replaced by the arguments' cell addresses."""
        tokens, env = _def_env(name, [xl_wrap(a) for a in args])
        def resolve(n):
            if n not in env:
                raise CalcError(f"{name}: unknown name {n}")
            return env[n]
        return xl_wrap(ExcelFormula(tokens, resolve).parse())
    def factor(self):
        kind, val = self.peek()
        if val in ("-", "+"):
            self.take(); f = self.factor()
            return f"-{f}" if val == "-" else f
        if val == "(":
            self.take(); v = self.expr(); self.take(")"); return f"({v})"
        if kind in ("num", "str", "lkp", "rng"):
            self.take(); return val
        if kind in ("ref", "xref"):
            self.take(); return self.resolve(val)
        if kind == "name":
            self.take()
            if self.peek()[1] == "(":
                if val in DEFS:
                    return self.call_def(val, self.args())
                if val not in EXCEL_FUNCS:
                    raise CalcError(f"No Excel form for function {val}")
                return xl_wrap(EXCEL_FUNCS[val](*[xl_wrap(a) for a in self.args()]))
            return self.resolve(val)                 # bare cell ref like D7
        raise CalcError(f"Unexpected token {val!r}")

# ---------------- where a row's fields sit in Excel ----------------
def id_col(row_id):
    """Column number of the cell a row id starts at: "B8" -> 2, "A13" -> 1 (no letters -> 2 = B)."""
    from openpyxl.utils import column_index_from_string
    m = re.match(r"\$?([A-Za-z]{1,3})\$?\d", str(row_id).strip())
    return column_index_from_string(m[1].upper()) if m else 2

def row_num(row_id):
    """Excel row of a row id: "A13" -> 13."""
    m = re.search(r"\d+", str(row_id))
    if not m:
        raise CalcError(f"Bad row id {row_id!r}")
    return int(m.group())

def shift_col(letter, by):
    """Column letter moved right by `by` columns: shift_col("B", -1) -> "A"."""
    from openpyxl.utils import column_index_from_string, get_column_letter
    return get_column_letter(max(1, column_index_from_string(letter) + by))

def fixed_cols(row_id):
    """Fixed layout, starting at the id's cell: "B8" -> description B, unit1 C, value1 D, value2 E,
    unit2 F, notes G; "A13" -> description A, unit1 B, value1 C, value2 D, unit2 E, notes F."""
    by = id_col(row_id) - 2
    return {k: shift_col(c, by) for k, c in {"description": "B", "unit1": "C", "unit": "C", "value1": "D",
            "value": "D", "value2": "E", "unit2": "F", "notes": "G"}.items()}

def row_cols(row):
    """Excel column of each field of a row (see fixed_cols)."""
    return fixed_cols(row["id"])

# ---------------- workbook evaluator ----------------
class Workbook:
    def __init__(self, tables, inputs=None):
        self.tables = {t["Table_id"]: t for t in tables}
        self.inputs = inputs or {}           # {"enquiry_sheet.abc": 500, "@TABLE.D11": 2257, ...}
        self.cache, self.stack, self.errors = {}, [], {}

    def _cell_key(self, table_id, cell):
        """'D7' -> (table_id, row id, 'value1'): the value of row 7 that sits in column D"""
        m = re.fullmatch(r"([A-Z]+)(\d+)", cell)
        if not m:
            raise CalcError(f"Bad cell {cell}")
        if table_id not in self.tables:
            raise CalcError(f"Unknown table {table_id}")
        row = self._find_row(table_id, f"B{m.group(2)}")
        cols = {c: k for k, c in row_cols(row).items() if k in VALUE_KEYS}
        if m.group(1) not in cols:
            raise CalcError(f"Can't map cell {cell} (row {m.group(2)} has values in {sorted(cols)})")
        return (table_id, row["id"], cols[m.group(1)])

    def _find_row(self, table_id, row_id):
        rows = self.tables[table_id]["rows"]
        for r in rows:
            if r["id"] == row_id:
                return r
        n = re.sub(r"\D", "", row_id)                 # same Excel row: "B11" finds the row "A11"
        for r in rows:
            if n and re.sub(r"\D", "", r["id"]) == n:
                return r
        raise CalcError(f"Row {row_id} not found in {table_id}")

    def value(self, table_id, row_id, key):
        if key in ("value1", "value"):             # single-value rows use "value" (column D)
            row = self._find_row(table_id, row_id)
            if key not in row:
                key = "value" if key == "value1" else "value1"
        k = (table_id, row_id, key)
        if k in self.cache:
            return self.cache[k]
        if k in self.stack:
            chain = " -> ".join(f"{t}.{r}.{c}" for t, r, c in self.stack + [k])
            raise CalcError(f"Circular reference: {chain}")
        # user-supplied override for this cell (e.g. the '???' inputs): "@T1.B9.value1" or "@T1.D9"
        try:
            col = row_cols(self._find_row(table_id, row_id)).get(key, "")
        except CalcError:
            col = ""
        for override in (f"{table_id}.{row_id}.{key}", f"{table_id}.{col}{row_id[1:]}"):
            if override in self.inputs:
                self.cache[k] = self.inputs[override]
                return self.cache[k]

        self.stack.append(k)
        try:
            raw = self._find_row(table_id, row_id).get(key)
            self.cache[k] = self._eval_raw(raw, table_id)
            return self.cache[k]
        finally:
            self.stack.pop()

    def _eval_raw(self, raw, table_id):
        if raw is None or raw in ("-", ""):
            return None
        if isinstance(raw, (int, float)):
            return raw
        if raw.strip() == "???":
            raise CalcError("Missing input ('???') - supply it via inputs")
        if not raw.startswith("="):
            try:
                return float(raw)
            except ValueError:
                return raw                          # plain text
        return Parser(tokenize(raw[1:]), lambda ref: self._resolve(ref, table_id)).parse()

    def _resolve(self, ref, current_table):
        # allow table refs written without '@' (T1.D8 -> @T1.D8) unless it is an external input
        if not ref.startswith("@") and ref not in self.inputs and "@" + ref.split(".")[0] in self.tables:
            ref = "@" + ref
        if ref.startswith("@"):
            parts = ref.split(".")
            tid = parts[0]
            if tid not in self.tables:
                raise CalcError(f"Unknown table {tid}")
            if len(parts) == 2:                       # @T.D7
                return self._need_num(self.value(*self._cell_key(tid, parts[1])), ref)
            if len(parts) == 3:                       # @T.value1.B7 or @T.B7.value1
                key, row = (parts[1], parts[2]) if parts[1] in VALUE_KEYS else (parts[2], parts[1])
                if key not in VALUE_KEYS:
                    raise CalcError(f"Bad reference {ref} (expected one of {VALUE_KEYS})")
                return self._need_num(self.value(tid, row, key), ref)
        elif "." in ref:                              # external: enquiry_sheet.abc
            if ref not in self.inputs:
                raise CalcError(f"External input '{ref}' not supplied")
            return self.inputs[ref]
        else:                                         # bare D7 -> current table
            return self._need_num(self.value(*self._cell_key(current_table, ref)), ref)
        raise CalcError(f"Bad reference {ref}")

    @staticmethod
    def _need_num(v, ref):
        if v is None:
            raise CalcError(f"{ref} is empty")
        return v

    def evaluate_all(self):
        """Compute every value cell; store result in row['computed'][key]. Errors are collected, not fatal."""
        for tid, t in self.tables.items():
            for r in t["rows"]:
                r["computed"] = {}
                for key in VALUE_KEYS:
                    if key not in r:
                        continue
                    self.stack = []
                    try:
                        v = self.value(tid, r["id"], key)
                    except (CalcError, ZeroDivisionError, TypeError) as e:
                        v = None
                        self.errors[f"{tid}.{r['id']}.{key}"] = str(e)
                    r["computed"][key] = round(v, 6) if isinstance(v, float) else v
        return self

    # ---------- storage ----------
    def save_json(self, path):
        Path(path).write_text(json.dumps(list(self.tables.values()), indent=2), encoding="utf-8")

    def save_csv(self, path):
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["table_id", "row_id", "description", "value1", "unit1", "value2", "unit2", "error"])
            for tid, t in self.tables.items():
                for r in t["rows"]:
                    c = r.get("computed", {})
                    err = "; ".join(v for k, v in self.errors.items() if k.startswith(f"{tid}.{r['id']}."))
                    w.writerow([tid, r["id"], r.get("description"), c.get("value1", c.get("value")), r.get("unit1", r.get("unit")),
                                c.get("value2"), r.get("unit2"), err])

    def save_xlsx(self, path):
        """Excel workbook with live formulas. One worksheet per sheet (its tables together); each row sits on its original
        Excel row (id "B9" -> row 9) in columns B..G, so D9 here is the cell the formulas refer to.
        Enquiry values go on an 'Enquiry' sheet; sheets read via 'SHEET'!D9 are added with their
        values. Changing any number in Excel recalculates every cell that depends on it."""
        from openpyxl import Workbook as XlWorkbook
        from openpyxl.styles import Font, PatternFill
        book = XlWorkbook()
        book.remove(book.active)
        bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
        input_fill, err_font = PatternFill("solid", fgColor="FFF2CC"), Font(color="C00000")
        xl_col = KEY_COL                                         # value1/value -> D, value2 -> E
        norm = lambda s: re.sub(r"[^a-z0-9]", "", s.lower())

        # ---- lay out sheets: this calc's tables ("") + linked sheets (values only).
        # Tables of the same sheet share one worksheet, each on its own original rows. A table without
        # sheet_name belongs to the sheet named before it. A table whose rows are already used by
        # another table of that sheet gets a worksheet of its own, so no value is overwritten.
        groups = [("", list(self.tables.values()))] + list(getattr(self, "linked", {}).items())
        layout = {}                                              # (group, table id) -> (worksheet, row shift)
        data_rows = {}                                           # worksheet title -> rows holding table data
        def new_sheet(name):
            name = re.sub(r"[\\/*?:\[\]]", "", name)[:31] or "Sheet"
            while name in book.sheetnames:
                name = name[:28] + f"_{len(book.sheetnames)}"
            data_rows[name] = set()
            return book.create_sheet(name)
        row_nums = lambda t: [int(r["id"][1:]) for r in t["rows"]]
        for g, tables in groups:
            by_sheet, sheet = {}, None
            for t in tables:
                sheet = t.get("sheet_name") or sheet or t["Table_id"]
                by_sheet.setdefault(sheet, []).append(t)
            for sheet, ts in by_sheet.items():
                shift = max(0, 4 - min((n for t in ts for n in row_nums(t)), default=4))  # room for titles
                ws = new_sheet(sheet)
                for t in ts:
                    target, t_shift = ws, shift
                    rows = {n + shift for n in row_nums(t)}
                    if rows & data_rows[ws.title]:               # rows taken by another table
                        target = new_sheet(f"{sheet} {t['Table_id']}")
                        t_shift = max(0, 4 - min(row_nums(t), default=4))
                        rows = {n + t_shift for n in row_nums(t)}
                    data_rows[target.title] |= rows
                    layout[(g, t["Table_id"])] = (target, t_shift)

        enq = book.create_sheet("Enquiry")
        enq["A1"], enq["B1"] = "Input", "Value"
        enq["A1"].font = enq["B1"].font = bold
        enq_row = {}
        def enquiry_cell(name):
            if name not in enq_row:
                n = enq_row[name] = len(enq_row) + 2
                enq[f"A{n}"] = name
                if name in self.inputs:
                    enq[f"B{n}"] = self.inputs[name]
                else:
                    enq[f"B{n}"].fill = input_fill                   # not supplied yet - fill it in
            return f"Enquiry!$B${enq_row[name]}"
        for name in self.inputs:
            if not name.startswith("@"):
                enquiry_cell(name)

        # ---- lookup tables, each on its own sheet at its original rows (B17..), for VLOOKUP
        from openpyxl.utils import get_column_letter
        LOOKUP_XL.clear()
        for lname, lt in LOOKUPS.items():
            name = re.sub(r"[\\/*?:\[\]]", "", lt.get("Vlookup_table_name") or lt["Vlookup_table_id"])[:31]
            while name in book.sheetnames:
                name = name[:28] + f"_{len(book.sheetnames)}"
            lws = book.create_sheet(name)
            ncol = len(lt.get("columns", []))
            lrows = [(int(r["id"][1:]), r) for r in lt["rows"]]
            first = min((n for n, _ in lrows), default=3)
            shift = max(0, 3 - first)
            lws["B1"], lws["B1"].font = f"{name} ({lname})", Font(bold=True, size=12)
            for i, c in enumerate(lt["columns"], 1):
                h = lws.cell(row=first + shift - 1, column=i + 1, value=c.get("display_name", c["key"]))
                h.font, h.fill = bold, head_fill
                lws.column_dimensions[get_column_letter(i + 1)].width = 16
            for n, r in lrows:
                for i in range(1, ncol + 1):
                    lws.cell(row=n + shift, column=i + 1, value=lookup_cell(lt, r, i))
            last = max((n for n, _ in lrows), default=first) + shift
            LOOKUP_XL[lname] = (name, shift)

        # ---- reference -> Excel address
        tables_of = {(g, t["Table_id"]): t for g, ts in groups for t in ts}

        def row_of(g, tid, row_id):
            """Row by id, or by Excel row number ("B11" finds "A11")."""
            rows = tables_of[(g, tid)]["rows"]
            n = re.sub(r"\D", "", row_id)
            hit = next((r for r in rows if r["id"] == row_id), None) or next(
                (r for r in rows if n and re.sub(r"\D", "", r["id"]) == n), None)
            if hit is None:
                raise CalcError(f"Row {row_id} not found in {tid}")
            return hit

        def addr(g, tid, row_id, key, here):
            if (g, tid) not in layout:
                raise CalcError(f"Unknown table {tid}")
            ws, shift = layout[(g, tid)]
            r = row_of(g, tid, row_id)
            if key not in r and key in ("value1", "value"):   # single-value rows use "value"
                key = "value" if key == "value1" else "value1"
            cell = f"{row_cols(r)[key]}{row_num(r['id']) + shift}"
            return cell if ws is here else "'" + ws.title.replace("'", "''") + "'!" + cell

        def row_key(parts, ref, g, tid):
            if len(parts) == 1:
                m = re.fullmatch(r"([A-Z]+)(\d+)", parts[0].replace("$", ""))
                if not m:
                    raise CalcError(f"Bad reference {ref}")
                r = row_of(g, tid, f"B{m.group(2)}")
                cols = {c: k for k, c in row_cols(r).items() if k in VALUE_KEYS}
                if m.group(1) not in cols:
                    raise CalcError(f"Can't map cell {parts[0]} (row {m.group(2)} has values in {sorted(cols)})")
                return r["id"], cols[m.group(1)]
            if len(parts) == 2:
                key, row_id = (parts[0], parts[1]) if parts[0] in VALUE_KEYS else (parts[1], parts[0])
                if key in VALUE_KEYS:
                    return row_id, key
            raise CalcError(f"Bad reference {ref}")

        def xl_ref(ref, g, tid, here):
            if "!" in ref:                                   # 'THERMAL - VAP'!D9 / !@T1.B9.value1
                sheet, target = ref.rsplit("!", 1)
                g2 = norm(sheet.strip("'"))
                tids = [t2 for (gg, t2) in layout if gg == g2]
                if not tids:
                    raise CalcError(f"Sheet {sheet} not loaded")
                parts = target.replace("$", "").lstrip("@").split(".")
                if len(parts) > 1 and "@" + parts[0] in tids:
                    tids, parts = ["@" + parts[0]], parts[1:]
                row_id, key = row_key(parts, ref, g2, tids[0])
                tid2 = next((t2 for t2 in tids for r in dict(groups)[g2] if r["Table_id"] == t2
                             and any(x["id"] == row_id for x in r["rows"])), tids[0])
                return addr(g2, tid2, row_id, key, here)
            parts = ref.split(".")
            if not ref.startswith("@") and ref not in self.inputs and (g, "@" + parts[0]) in layout:
                parts[0] = "@" + parts[0]                    # T1.D8 -> @T1.D8
            if parts[0].startswith("@"):
                return addr(g, parts[0], *row_key(parts[1:], ref, g, parts[0]), here)
            if "." in ref:                                   # enquiry_sheet.abc
                return enquiry_cell(ref)
            return addr(g, tid, *row_key([ref], ref, g, tid), here)  # bare D7

        # ---- write the rows
        for g, tables in groups:
            for t in tables:
                tid = t["Table_id"]
                ws, shift = layout[(g, tid)]
                if ws["B1"].value is None:           # sheet title
                    ws["B1"] = ws.title + ("  - values from linked sheet" if g else "")
                    ws["B1"].font = Font(bold=True, size=13)
                # table title + header just above the table's first row, only on rows nothing else uses
                first = min((int(r["id"][1:]) for r in t["rows"]), default=4) + shift
                heads = {} if t.get("columns") else {"B": "Description", "C": "Unit", "D": "Value1",
                                                        "E": "Value2", "F": "Unit2"}
                for c in t.get("columns", []):
                    col = HEADER_COL.get(c["key"])
                    if col:
                        heads[col] = c.get("display_name", c["key"])
                heads["G"] = "Notes / errors (last Python run)"
                by = id_col(min(t["rows"], key=lambda r: row_num(r["id"]))["id"]) - 2 if t["rows"] else 0
                heads = {shift_col(c, by): h for c, h in heads.items()}
                title = {shift_col("B", by): f"{t.get('Table_name', '')} ({tid})"}
                for line, cells in ((first - 2, title), (first - 1, heads)):
                    if line < 2 or line in data_rows[ws.title] or any(ws[f"{c}{line}"].value for c in cells):
                        continue
                    for col, h in cells.items():
                        c = ws[f"{col}{line}"]
                        c.value, c.font = h, bold
                        if line == first - 1:
                            c.fill = head_fill
                for r in t["rows"]:
                    line = row_num(r["id"]) + shift
                    cols = row_cols(r)
                    for key in ("description", "unit1", "unit", "unit2"):
                        if key in r:
                            ws[f"{cols[key]}{line}"] = r[key]
                    notes = [] if g else [v for k, v in self.errors.items() if k.startswith(f"{tid}.{r['id']}.")]
                    for key in VALUE_KEYS:
                        if key not in r:                     # value1/value2 or value - never both in D
                            continue
                        cell = ws[f"{cols[key]}{line}"]
                        raw = r.get(key)
                        if g:                                # linked sheet: plain values
                            cell.value = r.get("computed", {}).get(key)
                        elif isinstance(raw, str) and raw.startswith("="):
                            try:
                                f = ExcelFormula(tokenize(raw[1:]), lambda ref: xl_ref(ref, g, tid, ws)).parse()
                                cell.value = "=" + xl_unwrap(f)
                            except (CalcError, TypeError) as e:
                                notes.append(f"{key}: formula not converted ({e})")
                        elif isinstance(raw, str) and raw.strip() == "???":
                            cell.fill = input_fill           # missing input - type it in Excel
                        elif raw not in (None, "-", ""):
                            try:
                                cell.value = float(raw) if isinstance(raw, str) else raw
                            except ValueError:
                                cell.value = raw
                        if not isinstance(cell.value, str) or cell.value.startswith("="):
                            cell.number_format = "#,##0.000"
                    if notes or r.get("notes"):
                        ws[f"{cols['notes']}{line}"] = "; ".join(notes) or r.get("notes")
                        if notes:
                            ws[f"{cols['notes']}{line}"].font = err_font
                for col, w in zip("ABCDEFG", [3, 36, 12, 16, 16, 12, 60]):
                    ws.column_dimensions[col].width = w
        enq.column_dimensions["A"].width, enq.column_dimensions["B"].width = 40, 20
        book.calculation.fullCalcOnLoad = True               # Excel computes every formula on open
        book.save(path)

    def report(self):
        for tid, t in self.tables.items():
            print(f"\n== {t['Table_name']} ({tid}) ==")
            for r in t["rows"]:
                c = r.get("computed", {})
                fmt = lambda v: f"{v:,.3f}" if isinstance(v, (int, float)) else str(v if v is not None else "-")
                print(f"{r['id']:<4} {r['description']:<36} {fmt(c.get('value1', c.get('value'))):>16} {r.get('unit1', r.get('unit', '')):<7}"
                      f" {fmt(c.get('value2')):>16} {r.get('unit2','')}")
        if self.errors:
            print("\nErrors:")
            for k, v in self.errors.items():
                print(f"  {k}: {v}")


NUM_WITH_UNIT_RE = re.compile(r"^\s*([-+]?\d+\.?\d*(?:[eE][-+]?\d+)?)\s*[A-Za-z(/%]")

def load_enquiry(path):
    """Read an enquiry JSON ({"fields": {name: {"value": ...}}}) into
    {"enquiry_sheet.<name>": value}. Numbers with units ("1000 kg/hr") become floats;
    fields with a null value are left out so they report as 'not supplied'."""
    fields = json.loads(Path(path).read_text(encoding="utf-8")).get("fields", {})
    inputs = {}
    for name, f in fields.items():
        v = f.get("value") if isinstance(f, dict) else f
        if v is None:
            continue
        if isinstance(v, str):
            s = v.strip()
            try:
                v = float(s)
            except ValueError:
                m = NUM_WITH_UNIT_RE.match(s)
                if m:
                    v = float(m.group(1))
        inputs[f"enquiry_sheet.{name}"] = v
    return inputs


def run(json_path, enquiry_path=None, out_path=None):
    """Evaluate one calc JSON (optionally with an enquiry sheet), print the report and
    save <out_path>.json / .csv (default: <json stem>_computed in the current folder)."""
    json_path = Path(json_path)
    tables = json.loads(json_path.read_text(encoding="utf-8"))
    inputs = load_enquiry(enquiry_path) if enquiry_path else {}
    load_lookups(json_path.parent)                 # lookup table JSONs next to the calc JSON
    defs = json_path.with_name(f"{json_path.stem} def.json")
    if defs.is_file():                             # f_def formulas: "sample thermal vap def.json"
        load_defs(defs)
    wb = Workbook(tables, inputs).evaluate_all()
    wb.report()
    out = Path(out_path) if out_path else Path(json_path.stem + "_computed")
    wb.save_json(f"{out}.json")
    wb.save_csv(f"{out}.csv")
    wb.save_xlsx(f"{out}.xlsx")
    print(f"\nSaved {out.name}.json, {out.name}.csv and {out.name}.xlsx")
    return wb


if __name__ == "__main__":
    import sys
    HERE = Path(__file__).resolve().parent
    # optional overrides: python thermal_vap_calc_engine.py [calc json] [enquiry json]
    CALC_JSON = sys.argv[1] if len(sys.argv) > 1 else HERE.parent / "json" / "sample thermal vap.json"
    ENQUIRY_JSON = sys.argv[2] if len(sys.argv) > 2 else HERE.parent / "enquiry" / "cl2 enquiry.json"
    OUTPUT = HERE / "thermal_vap_computed"          # -> thermal_vap_computed.json / .csv
    # constants are written directly in the calc JSON as plain numbers
    run(CALC_JSON, ENQUIRY_JSON, OUTPUT)
