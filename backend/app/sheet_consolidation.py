"""
Excel sheet consolidation.

For a workbook attachment, reads every sheet with openpyxl and writes a
sparse "raw" JSON file per sheet describing each non-empty cell: its
coordinate, row/col, cached value, formula text (if any), and merge range
(if it's the anchor of a merged region). This raw, cell-level view is kept
separate from the pandas-based extract_text_from_excel() in extraction.py
so both can be compared/used independently.

Adapted from the standalone consolidate_data_from_sheets.py script.
"""

import json
import re
from pathlib import Path

import openpyxl


def safe_filename(name: str) -> str:
    """Make a sheet name filesystem-safe for use in an output filename."""
    return re.sub(r"[^A-Za-z0-9_\-]+", "_", name).strip("_") or "sheet"


def extract_sheet(ws_values, ws_formulas) -> list:
    """
    ws_values:   worksheet loaded with data_only=True  (cached computed values)
    ws_formulas: worksheet loaded with data_only=False (raw formula text)

    Returns a list of cell dicts, ordered by row then column, containing
    only non-empty cells (and only merge-anchor cells within merged
    regions).
    """
    # Map anchor cell coord -> merge range string, and the full set of
    # coordinates that belong to some merged region (anchor + members).
    merge_map = {}
    merged_member_cells = set()
    for mrange in ws_values.merged_cells.ranges:
        anchor = mrange.coord.split(":")[0]
        merge_map[anchor] = str(mrange)
        for row in ws_values.iter_rows(
            min_row=mrange.min_row, max_row=mrange.max_row,
            min_col=mrange.min_col, max_col=mrange.max_col,
        ):
            for cell in row:
                merged_member_cells.add(cell.coordinate)

    cells = []
    for row_v, row_f in zip(ws_values.iter_rows(), ws_formulas.iter_rows()):
        for cell_v, cell_f in zip(row_v, row_f):
            coord = cell_v.coordinate

            # Skip non-anchor cells inside a merged region — they hold no
            # data of their own; the anchor cell carries the value.
            if coord in merged_member_cells and coord not in merge_map:
                continue

            is_formula = isinstance(cell_f.value, str) and cell_f.value.startswith("=")
            value = cell_v.value  # cached / computed value
            formula = cell_f.value if is_formula else None

            # Skip genuinely empty cells (no value, no formula).
            if value is None and formula is None:
                continue

            entry = {
                "cell": coord,
                "row": cell_v.row,
                "col": cell_v.column,
                "value": value,
            }
            if formula is not None:
                entry["formula"] = formula
                if value is None:
                    # Formula present but no cached value was saved in the
                    # file (e.g. workbook wasn't recalculated before save).
                    entry["value_stale_or_missing"] = True
            if coord in merge_map:
                entry["merge_range"] = merge_map[coord]

            cells.append(entry)

    return cells


def consolidate_workbook_sheets(xlsx_path: str, out_dir: str, filename_suffix: str = "") -> list:
    """Read xlsx_path and write one <sheetname>_raw<filename_suffix>.json
    file per sheet into out_dir. filename_suffix lets the caller embed a
    timestamp/thread id in the filename (e.g. "-17092026_180254-<thread_id>")
    without this module needing to know about either. Returns the list of
    written file paths."""
    wb_values = openpyxl.load_workbook(xlsx_path, data_only=True)
    wb_formulas = openpyxl.load_workbook(xlsx_path, data_only=False)

    out_path = Path(out_dir)
    out_path.mkdir(parents=True, exist_ok=True)
    written = []

    for sheetname in wb_values.sheetnames:
        cells = extract_sheet(wb_values[sheetname], wb_formulas[sheetname])
        out_file = out_path / f"{safe_filename(sheetname)}_raw{filename_suffix}.json"
        out_file.write_text(
            json.dumps({"sheet_name": sheetname, "cells": cells}, indent=2, ensure_ascii=False, default=str),
            encoding="utf-8",
        )
        written.append(str(out_file))

    return written
