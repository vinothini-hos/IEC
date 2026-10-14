"""
Design calculation sheets (Thermal Vap, Steam, Wall Thickness, ...) run
against one extracted equipment item.

Each sheet has its own calc engine script, holding that sheet's f_def
functions (design_calc_inputs/engines/<sheet>_calc_engine.py), in
settings.design_calc_engine_dir, and its calc JSON (plus any VLOOKUP table JSONs)
in settings.design_calc_json_dir (backend/design_calc_inputs/json). The
engines are loaded as-is - they've diverged (f_def2/f_def3 mean different
things per sheet, wall thick has merged cells and extra columns) so they
aren't merged into one module.
The equipment item's fields are the enquiry sheet (enquiry_sheet.<field_key>),
exactly the {"fields": {key: {"value": ...}}} shape the engines' load_enquiry
reads.

To add a sheet: drop its engine into design_calc_engine_dir, its calc/VLOOKUP JSONs
into design_calc_json_dir, and add an entry to SHEETS. `defs` (optional) names
the JSON holding that sheet's f_def formulas, for engines with load_defs. `linked` lists sheets
whose computed values it reads ('THERMAL - VAP'!D9) - those run first and are
passed in as `linked`.
"""
import importlib.util
import json
import math
import re
import threading
import uuid
from datetime import datetime
from pathlib import Path

from .config import settings

SHEETS = [
    {
        "key": "thermal_vap",
        "label": "Thermal Vap",
        "engine": "thermal_vap_calc_engine.py",
        "calc_json": "sample thermal vap.json",
        "defs": "sample thermal vap def.json",
        "linked": [],
    },
    {
        "key": "piping",
        "label": "Piping",
        "engine": "piping_calc_engine.py",
        "calc_json": "sample piping json.json",
        "defs": "sample piping def.json",
        "linked": [],
    },
    {
        "key": "steam",
        "label": "Steam",
        "engine": "steam_calc_engine.py",
        "calc_json": "sample steam json.json",
        "defs": "sample steam def.json",
        "linked": ["thermal_vap"],
    },
    {
        "key": "vap_press_dens",
        "label": "Vap Press / Density",
        "engine": "vap_press_dens_calc_engine.py",
        "calc_json": "sample vap press-dens json.json",
        "defs": "sample vap press-dens def.json",
        "linked": [],
    },
    {
        "key": "wall_thick",
        "label": "Wall Thickness",
        "engine": "wall_thick_calc_engine.py",
        "calc_json": "sample wall thick json.json",
        "defs": "sample wall thick def.json",
        "linked": [],
    },
]
SHEETS_BY_KEY = {s["key"]: s for s in SHEETS}

# The engines keep lookup tables in module globals (LOOKUPS / LOOKUP_XL),
# so only one run at a time.
_run_lock = threading.Lock()
_engines = {}  # sheet key -> (engine file mtime, module)


def _load_engine(sheet):
    """Import the sheet's engine script; re-imported when the file changes so
    edits are picked up without restarting the server."""
    path = Path(settings.design_calc_engine_dir) / sheet["engine"]
    mtime = path.stat().st_mtime
    cached = _engines.get(sheet["key"])
    if cached and cached[0] == mtime:
        return cached[1]
    spec = importlib.util.spec_from_file_location(f"design_calc_engine_{sheet['key']}", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    _engines[sheet["key"]] = (mtime, module)
    return module


def _sheet_key(name):
    """Same loose sheet-name match as the engines' sheet_key: 'THERMAL - VAP' == 'Thermal vap'."""
    return "".join(ch for ch in name.lower() if ch.isalnum())


def _linked_tables(sheet, computed):
    """{sheet_key(sheet_name): [computed tables]} for the sheets this one reads from.
    A table without its own sheet_name belongs to the sheet named before it."""
    linked = {}
    for dep in sheet["linked"]:
        name = SHEETS_BY_KEY[dep]["label"]
        for t in computed.get(dep) or []:
            name = t.get("sheet_name") or name
            linked.setdefault(_sheet_key(name), []).append(t)
    return linked


def _json_safe(v):
    """inf/nan (e.g. from a LN of a huge number) aren't valid JSON."""
    if isinstance(v, float) and not math.isfinite(v):
        return None
    if isinstance(v, dict):
        return {k: _json_safe(x) for k, x in v.items()}
    if isinstance(v, list):
        return [_json_safe(x) for x in v]
    return v


def _enquiry_inputs(engine, enquiry_path):
    """The engine's enquiry inputs ("enquiry_sheet.<field>" = value, numbers
    with units parsed) plus each field's attributes as it's laid out in the
    extraction JSON, so a calc JSON can also write
    =enquiry_sheet.<field>.value / .status / .label / .source."""
    inputs = engine.load_enquiry(enquiry_path)
    fields = json.loads(Path(enquiry_path).read_text(encoding="utf-8")).get("fields", {})
    for name, field in fields.items():
        if not isinstance(field, dict):
            continue
        for attr, v in field.items():
            if attr != "value" and v is not None:
                inputs[f"enquiry_sheet.{name}.{attr}"] = v
        if f"enquiry_sheet.{name}" in inputs:
            inputs[f"enquiry_sheet.{name}.value"] = inputs[f"enquiry_sheet.{name}"]
    return inputs


def _apply_overrides(tables, cells):
    """Put the user's edited values ({table_id: {row_id: {key: value}}}) in
    place of the calc JSON's formulas/constants before evaluating, so they're
    used by every dependent cell and written to the .xlsx as plain values.
    Returns {(table_id, row_id): {key: original value}} for the edited rows."""
    originals = {}
    for t in tables:
        rows = cells.get(t.get("Table_id"), {})
        for r in t.get("rows", []):
            for key, value in rows.get(r.get("id"), {}).items():
                originals.setdefault((t["Table_id"], r["id"]), {})[key] = r.get(key)
                r[key] = value
    return originals


def _run_sheet(sheet, enquiry_path, computed, out_dir, overrides):
    engine = _load_engine(sheet)
    calc_path = Path(settings.design_calc_json_dir) / sheet["calc_json"]
    tables = json.loads(
        calc_path.read_text(encoding="utf-8"),
        object_pairs_hook=getattr(engine, "read_pairs", None),  # wall thick: "merged" keys
    )
    inputs = _enquiry_inputs(engine, enquiry_path)
    engine.LOOKUPS.clear()
    engine.load_lookups(calc_path.parent)
    if sheet.get("defs"):                   # f_def formulas from the sheet's def JSON
        engine.load_defs(calc_path.parent / sheet["defs"])

    if sheet["linked"]:
        wb = engine.Workbook(tables, inputs, _linked_tables(sheet, computed))
    else:
        wb = engine.Workbook(tables, inputs)
    originals = _apply_overrides(wb.tables.values(), overrides.get(sheet["key"], {}))
    wb.evaluate_all()
    wb.save_xlsx(out_dir / f"{sheet['key']}.xlsx")
    # attached only after save_xlsx - wall thick treats unknown row keys as extra columns
    for t in wb.tables.values():
        for r in t["rows"]:
            if (t["Table_id"], r["id"]) in originals:
                r["edited"] = originals[(t["Table_id"], r["id"])]
    lookup_sheets = {name for name, _ in getattr(engine, "LOOKUP_XL", {}).values()}
    return list(wb.tables.values()), wb.errors, lookup_sheets


LINKED_SHEET_MARK = "values from linked sheet"   # engines' title on linked-sheet value copies
ENQUIRY_SHEET = "Enquiry"
# 'Sheet name'!$B$4 or Sheet!B4:C9 - a reference into another worksheet
_SHEET_REF_RE = re.compile(r"('(?:[^']|'')+'|[A-Za-z_][\w.]*)!(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?")


def _quote_sheet(name):
    """VLOOKUP1 stays bare, as Excel writes it; 'Thermal vap' needs quotes."""
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name) and not re.fullmatch(r"[A-Za-z]{1,3}\d+", name):
        return name
    return "'" + name.replace("'", "''") + "'"


def _xl_literal(v, blank):
    """A cell value written inline in a formula: 12.5, "text", TRUE."""
    if v is None:
        return blank
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        text = str(int(v)) if float(v).is_integer() else repr(float(v))
        return f"({text})" if v < 0 else text
    return '"' + str(v).replace('"', '""') + '"'


def _plain_value(formula):
    """A formula left with just a value (an inlined enquiry input: =1000,
    ="CL2 Vaporizer") is written as that plain value."""
    m = re.fullmatch(r"=\(?(-?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?)\)?", formula)
    if m:
        n = float(m[1])
        return int(n) if n.is_integer() else n
    m = re.fullmatch(r'="((?:[^"]|"")*)"', formula)
    return m[1].replace('""', '"') if m else formula


def _same_values(a, b):
    values = lambda ws: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
    return values(a) == values(b)


def _copy_cell(src, dst, value):
    from copy import copy
    dst.value = value
    if src.has_style:
        dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
        dst.alignment, dst.protection = copy(src.alignment), copy(src.protection)
        dst.number_format = src.number_format
    if src.comment:
        dst.comment = copy(src.comment)


def _copy_sheet(src, dst, fix_formula):
    for row in src.iter_rows():
        for c in row:
            if c.value is None and not c.has_style:
                continue
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                v = _plain_value(fix_formula(v))
            _copy_cell(c, dst.cell(row=c.row, column=c.column), v)
    for rng in src.merged_cells.ranges:
        dst.merge_cells(str(rng))
    for key, dim in src.column_dimensions.items():
        dst.column_dimensions[key].width = dim.width


def _combine_workbooks(paths, out_path, lookup_sheets=()):
    """Merge the engines' per-sheet workbooks into one, keeping formulas live:
    - the Enquiry sheet is left out; a reference to an enquiry input becomes
      its value ('Enquiry'!$B$45 -> 1000);
    - the VLOOKUP table sheets (`lookup_sheets`) are kept, hidden, after all the
      calc sheets, so lookups stay as ranges: VLOOKUP(D37,VLOOKUP1!$B$36:$C$63,2,FALSE);
    - each engine's linked-sheet value copy (steam's "Thermal vap" values) and
      any VLOOKUP sheet identical to one already taken are dropped, so their
      references point at the sheet already in the combined book;
    - any other clash of names is renamed "<name> (2)" and references to it
      rewritten."""
    from openpyxl import Workbook, load_workbook
    book = Workbook()
    book.remove(book.active)
    planned = {}                            # combined-book title -> source sheet
    calc_todo, lookup_todo = [], []

    for path in paths:
        src = load_workbook(path)           # formulas come back as strings
        rename = {}

        def fix_formula(f, src=src, rename=rename):
            def sheet_ref(m):
                name = m[1][1:-1].replace("''", "'") if m[1].startswith("'") else m[1]
                first, last = m[2].replace("$", ""), m[3] and m[3].replace("$", "")
                if name == ENQUIRY_SHEET and name in src.sheetnames and not last:
                    return _xl_literal(src[name][first].value, "0")   # a blank cell counts as 0, as in Excel
                return _quote_sheet(rename.get(name, name)) + "!" + m[2] + (f":{m[3]}" if m[3] else "")
            return _SHEET_REF_RE.sub(sheet_ref, f)

        for ws in src.worksheets:
            if ws.title == ENQUIRY_SHEET:
                continue
            taken = planned.get(ws.title)
            if taken is not None and (
                    LINKED_SHEET_MARK in str(ws["B1"].value or "") or _same_values(ws, taken)):
                continue                    # use the sheet already in the combined book
            title, n = ws.title, 2
            while title in planned:
                title = f"{ws.title[:26]} ({n})"
                n += 1
            planned[title] = ws
            rename[ws.title] = title
            todo = lookup_todo if ws.title in lookup_sheets else calc_todo
            todo.append((ws, title, fix_formula))

    for ws, title, fix_formula in calc_todo + lookup_todo:
        _copy_sheet(ws, book.create_sheet(title), fix_formula)
    for _, title, _ in lookup_todo if calc_todo else []:   # Excel needs one visible sheet
        book[title].sheet_state = "hidden"  # still there for the ranges; Unhide to inspect

    if not book.sheetnames:
        book.create_sheet("Design calc")
    book.calculation.fullCalcOnLoad = True  # Excel computes every formula on open
    book.save(out_path)


def latest_extraction_file(thread_id: uuid.UUID, equipment_index: int):
    """Most recent data extraction JSON for this equipment item, or None.

    template_consolidation writes storage/data_extraction/{thread_id}/
    <idx>_<type>_<name>-<timestamp>-<thread_id>.json on every extraction and
    deletes the files of equipment dropped from the latest run, so the
    remaining <idx>_<type>_<name> prefixes, in order, line up with
    Thread.extraction_result["equipment_items"]. Newest by mtime, as the
    DDMMYYYY timestamp doesn't sort across month boundaries."""
    out_dir = Path(settings.data_extraction_storage_dir) / str(thread_id)
    by_prefix = {}
    for path in out_dir.glob(f"*-{thread_id}.json"):
        prefix = path.name[: -len(f"-{thread_id}.json")].rsplit("-", 1)[0]
        by_prefix.setdefault(prefix, []).append(path)
    prefixes = sorted(by_prefix)
    if not 0 <= equipment_index < len(prefixes):
        return None
    return max(by_prefix[prefixes[equipment_index]], key=lambda p: p.stat().st_mtime)


def run_for_equipment(thread_id: uuid.UUID, equipment_index: int, equipment_item: dict) -> dict:
    """Run every sheet (dependencies first, SHEETS order) against the equipment
    item, save result.json + one design_calc.xlsx with every sheet under
    storage, and return the result.
    The enquiry sheet input is the equipment's most recent data extraction
    JSON; `equipment_item` (from Thread.extraction_result) is only used when
    that file is gone. A sheet that fails outright is reported with `error`;
    the rest still run."""
    out_dir = Path(settings.design_calc_storage_dir) / str(thread_id) / str(equipment_index)
    out_dir.mkdir(parents=True, exist_ok=True)
    enquiry_path = latest_extraction_file(thread_id, equipment_index)
    if enquiry_path is None:
        enquiry_path = out_dir / "enquiry.json"
        enquiry_path.write_text(json.dumps(equipment_item, indent=2, default=str), encoding="utf-8")
    equipment_item = json.loads(enquiry_path.read_text(encoding="utf-8"))
    overrides = _load_overrides(out_dir)

    computed, sheets, lookup_sheets = {}, [], set()
    with _run_lock:
        for sheet in SHEETS:
            entry = {"key": sheet["key"], "label": sheet["label"]}
            (out_dir / f"{sheet['key']}.xlsx").unlink(missing_ok=True)
            try:
                tables, errors, lookups = _run_sheet(sheet, enquiry_path, computed, out_dir, overrides)
                lookup_sheets |= lookups
                computed[sheet["key"]] = tables
                entry.update(tables=tables, errors=errors)
            except Exception as e:
                entry.update(tables=[], errors={}, error=f"{type(e).__name__}: {e}")
            sheets.append(entry)

    combined = out_dir / "design_calc.xlsx"
    combined.unlink(missing_ok=True)
    sheet_files = [out_dir / f"{s['key']}.xlsx" for s in SHEETS]
    xlsx_error = None
    try:
        _combine_workbooks([p for p in sheet_files if p.is_file()], combined, lookup_sheets)
    except Exception as e:
        xlsx_error = f"Excel not generated: {type(e).__name__}: {e}"

    result = _json_safe({
        "equipment_index": equipment_index,
        "equipment_name": equipment_item.get("equipment_name"),
        "enquiry_file": enquiry_path.name,
        "ran_at": datetime.now().isoformat(timespec="seconds"),
        "sheets": sheets,
        "has_xlsx": combined.is_file(),
        "xlsx_error": xlsx_error,
        "pending_recalc": False,
    })
    (out_dir / "result.json").write_text(json.dumps(result, indent=2, default=str), encoding="utf-8")
    return result


def _inputs_mtime(thread_id: uuid.UUID, equipment_index: int) -> float:
    """Newest modification time of everything a run reads: engine scripts,
    calc/VLOOKUP JSONs, the equipment's latest data extraction file and the
    user's edited cells."""
    paths = [Path(settings.design_calc_engine_dir) / s["engine"] for s in SHEETS]
    paths += list(Path(settings.design_calc_json_dir).glob("*.json"))
    enquiry = latest_extraction_file(thread_id, equipment_index)
    if enquiry is not None:
        paths.append(enquiry)
    paths.append(Path(settings.design_calc_storage_dir) / str(thread_id) / str(equipment_index) / "overrides.json")
    return max((p.stat().st_mtime for p in paths if p.is_file()), default=0)


def has_result(thread_id) -> bool:
    """Any equipment item of this thread has a saved calculation."""
    return any((Path(settings.design_calc_storage_dir) / str(thread_id)).glob("*/result.json"))


def load_result(thread_id: uuid.UUID, equipment_index: int):
    """Last saved result, or None if there isn't one or it's stale - an
    engine, calc JSON or the extraction changed since it was calculated -
    so the page runs a fresh calculation instead of showing old values."""
    path = Path(settings.design_calc_storage_dir) / str(thread_id) / str(equipment_index) / "result.json"
    if not path.is_file() or path.stat().st_mtime < _inputs_mtime(thread_id, equipment_index):
        return None
    result = json.loads(path.read_text(encoding="utf-8"))
    # results saved before a SHEETS reorder still come back in SHEETS order
    order = {s["key"]: i for i, s in enumerate(SHEETS)}
    result["sheets"].sort(key=lambda e: order.get(e["key"], len(order)))
    return result


def xlsx_path(thread_id: uuid.UUID, equipment_index: int):
    """The combined workbook (all sheets) of the last run, or None."""
    path = Path(settings.design_calc_storage_dir) / str(thread_id) / str(equipment_index) / "design_calc.xlsx"
    return path if path.is_file() else None


def _load_overrides(out_dir: Path) -> dict:
    """User-edited cells: {sheet_key: {table_id: {row_id: {key: value}}}}."""
    path = out_dir / "overrides.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.is_file() else {}


def _parse_cell_value(value):
    """Numbers typed in the UI ("12.5") are stored as numbers, anything else as text."""
    if isinstance(value, str):
        try:
            return float(value.strip())
        except ValueError:
            return value
    return value


def set_cell(thread_id: uuid.UUID, equipment_index: int, sheet_key: str,
             table_id: str, row_id: str, key: str, value) -> dict:
    """Edit one cell of the last result, or revert it when `value` is None/"".
    The edit is saved to overrides.json (used by every later run) and written
    into result.json right away; cells that depend on it keep their old values
    until the next run, flagged by result["pending_recalc"]."""
    out_dir = Path(settings.design_calc_storage_dir) / str(thread_id) / str(equipment_index)
    result_path = out_dir / "result.json"
    if not result_path.is_file():
        raise ValueError("Run the calculation first")
    with _run_lock:
        result = json.loads(result_path.read_text(encoding="utf-8"))
        entry = next((e for e in result["sheets"] if e["key"] == sheet_key), None)
        table = next((t for t in (entry or {}).get("tables", []) if t.get("Table_id") == table_id), None)
        row = next((r for r in (table or {}).get("rows", []) if r.get("id") == row_id), None)
        if row is None:
            raise ValueError(f"No cell {sheet_key} {table_id} {row_id}")
        edited = row.get("edited", {})
        if key not in row.get("computed", {}) and key not in edited:
            raise ValueError(f"{row_id} has no calculated value {key!r}")

        overrides = _load_overrides(out_dir)
        cells = overrides.setdefault(sheet_key, {}).setdefault(table_id, {}).setdefault(row_id, {})
        if value is None or (isinstance(value, str) and not value.strip()):
            cells.pop(key, None)
            if key in edited:
                row[key] = edited.pop(key)
            row.setdefault("computed", {})[key] = None      # known again after recalculating
        else:
            value = _parse_cell_value(value)
            cells[key] = value
            edited.setdefault(key, row.get(key))
            row[key] = value
            row.setdefault("computed", {})[key] = value
        if edited:
            row["edited"] = edited
        else:
            row.pop("edited", None)
        entry.get("errors", {}).pop(f"{table_id}.{row_id}.{key}", None)
        result["pending_recalc"] = True

        # drop emptied levels so overrides.json only lists real edits
        if not cells:
            del overrides[sheet_key][table_id][row_id]
        if not overrides[sheet_key][table_id]:
            del overrides[sheet_key][table_id]
        if not overrides[sheet_key]:
            del overrides[sheet_key]
        (out_dir / "overrides.json").write_text(json.dumps(overrides, indent=2), encoding="utf-8")
        result = _json_safe(result)
        result_path.write_text(json.dumps(result, indent=2, default=str), encoding="utf-8")
    return result
