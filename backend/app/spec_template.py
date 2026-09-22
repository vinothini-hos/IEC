"""
Fills the blank SO2/CL2 query-sheet templates with the extracted spec values,
for attaching to the clarification email. Row layout must stay in sync with
SO2_TEMPLATE_FIELDS / CL2_TEMPLATE_FIELDS in template_consolidation.py — both
templates list the same 23 fields in the same order, starting at row 5, with
the value going in column E ("Client Input") — column D is "UOM" (a static
unit label already printed in the template, e.g. "kg/hr", and must not be
touched here).
"""

import io
import re

from .config import settings
from .template_consolidation import get_template_fields, has_real_value

TEMPLATE_PATHS = {
    "SO2": lambda: settings.so2_template_path,
    "CL2": lambda: settings.cl2_template_path,
}

FIRST_ROW = 5
VALUE_COLUMN = 5  # column E ("Client Input") — column D ("UOM") is left alone
CLIENT_NAME_CELL = "B1"
MISSING_FILL_COLOR = "FFC7CE"  # same pink Excel uses for its built-in "Bad" cell style


def _client_name(recipient: str) -> str:
    """"N VINOTHINI <nvinothin21@gmail.com>" -> "N VINOTHINI"."""
    if not recipient:
        return ""
    match = re.match(r'^"?([^"<]+)"?\s*<.*>$', recipient.strip())
    return match.group(1).strip() if match else recipient.strip()


def fill_template_workbook(equipment_name: str, type_label: str, fields: dict, recipient: str):
    """Returns (filename, bytes) for the filled workbook, or (None, None) if
    there's no template for this equipment type."""
    import openpyxl
    from openpyxl.styles import PatternFill

    path_fn = TEMPLATE_PATHS.get(type_label)
    if not path_fn:
        return None, None

    field_defs, _ = get_template_fields(equipment_name)
    if not field_defs:
        return None, None

    wb = openpyxl.load_workbook(path_fn())
    ws = wb.active

    ws[CLIENT_NAME_CELL] = _client_name(recipient)

    missing_fill = PatternFill(
        start_color=MISSING_FILL_COLOR, end_color=MISSING_FILL_COLOR, fill_type="solid"
    )

    for offset, (key, _label) in enumerate(field_defs):
        entry = fields.get(key) or {}
        value = entry.get("value")
        cell = ws.cell(row=FIRST_ROW + offset, column=VALUE_COLUMN, value=value if value else "-")
        if not has_real_value(entry):
            cell.fill = missing_fill

    buf = io.BytesIO()
    wb.save(buf)

    filename = f"IEC Queries - {type_label} template (filled).xlsx"
    return filename, buf.getvalue()
