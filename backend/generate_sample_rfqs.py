r"""
Generates realistic-but-incomplete sample RFQ documents (PDF/Word/Excel) for
SO2 and CL2 vaporizer enquiries, for testing the extraction pipeline against
varied real-world document formats and dummy customer data.

Also writes sample_data/{TYPE}/manifest.json — the ground-truth spec fields
and customer info used to generate each file. index_sample_data.py reads
that manifest to populate a Qdrant "sample" collection directly, without
re-running the (slow, local) LLM extraction on all 200 documents.

This only writes files under sample_data/ — it does not touch the database,
attachments storage, or the real RAG index.

Usage:
    python generate_sample_rfqs.py
"""

import json
import random
import time
from pathlib import Path

import openpyxl
from docx import Document
from fpdf import FPDF
from fpdf.enums import XPos, YPos

from app.template_consolidation import CL2_TEMPLATE_FIELDS, SO2_TEMPLATE_FIELDS

OUTPUT_DIR = Path(__file__).parent / "sample_data"
SAMPLES_PER_TYPE = 100
FORMAT_COUNTS = {"pdf": 40, "docx": 30, "xlsx": 30}  # must sum to SAMPLES_PER_TYPE

random.seed(42)  # reproducible output across runs

# ---------------------------------------------------------------------------
# Dummy customer pool
# ---------------------------------------------------------------------------

COMPANIES = [
    "Trident Industrial Gases", "Sarna Chemicals Pvt Ltd", "Vantage Aqua Treatment",
    "Meridian Pulp & Paper", "Coastal Water Works", "Bharat Agro Chemicals",
    "Nilgiri Sugar Mills", "Deccan Textiles Ltd", "Orient Petrochemicals",
    "Sundaram Effluent Systems", "Kaveri Fertilizers", "Anand Dyestuffs",
    "Shreeji Water Solutions", "Ganga Paper Mills", "Vindhya Chemicals",
    "Pioneer Bleaching Works", "Konkan Coast Desalination", "Amber Specialty Chemicals",
    "Northline Water Utility", "Suryoday Sugar Industries", "Greenfield Agro Processing",
    "Malwa Textile Chemicals", "Star Leather Industries", "Blue Nile Water Treatment",
    "Everest Pulp Corporation", "Rajdhani Municipal Waterworks", "Coromandel Process Chemicals",
    "Sagar Effluent Treatment", "Vishwakarma Steel & Chem", "Harbor City Utilities",
]

LOCATIONS = [
    "Vapi, Gujarat", "Raipur, Chhattisgarh", "Nashik, Maharashtra", "Salem, Tamil Nadu",
    "Vijayawada, Andhra Pradesh", "Panipat, Haryana", "Bhilai, Chhattisgarh",
    "Rourkela, Odisha", "Hosur, Tamil Nadu", "Ankleshwar, Gujarat", "Aurangabad, Maharashtra",
    "Kanpur, Uttar Pradesh", "Vizag, Andhra Pradesh", "Belgaum, Karnataka",
    "Jamshedpur, Jharkhand", "Bharuch, Gujarat", "Coimbatore, Tamil Nadu",
    "Indore, Madhya Pradesh", "Surat, Gujarat", "Durgapur, West Bengal",
]

CONTACT_NAMES = [
    "Ashok Menon", "Priya Raman", "Suresh Iyer", "Neha Kulkarni", "Rajesh Patil",
    "Kavita Nair", "Manoj Verma", "Deepa Krishnan", "Vikram Rao", "Anita Desai",
    "Sanjay Gupta", "Meera Pillai", "Arvind Shetty", "Lakshmi Narayan", "Ramesh Chandra",
    "Sunita Joshi", "Karthik Subramaniam", "Pooja Agarwal", "Vinod Bhatt", "Divya Menon",
]

TITLES = ["Procurement Manager", "Plant Engineer", "Purchase Head", "Technical Manager",
          "Project Engineer", "Works Manager", "General Manager - Operations"]


def _slug(text: str) -> str:
    return "".join(c.lower() if c.isalnum() else "" for c in text)[:16] or "customer"


def random_customer(index: int) -> dict:
    company = random.choice(COMPANIES)
    contact = random.choice(CONTACT_NAMES)
    return {
        "company": company,
        "location": random.choice(LOCATIONS),
        "contact_person": contact,
        "contact_title": random.choice(TITLES),
        "contact_email": f"{_slug(contact)}@{_slug(company)}.example.com",
        "contact_phone": f"+91 {random.randint(70000, 99999)}{random.randint(10000, 99999)}",
        "enquiry_reference": f"RFQ-{random.randint(2024, 2026)}-{1000 + index}",
    }


# ---------------------------------------------------------------------------
# Field value pools (per equipment type) — phrased like real spec sheets
# ---------------------------------------------------------------------------

COMMON_POOLS = {
    "qty": ["1", "2", "1 (with 1 standby)", "3"],
    "location_installation": ["Indoor", "Outdoor"],
    "system_bolt_down_location": [
        "Outdoor RCC foundation, dedicated vaporizer skid area",
        "Indoor plant room, near reactor building",
        "Existing skid area next to old vaporizer",
        "New RCC plinth near tonner shed",
        "Rooftop platform, structural steel skid",
        "Ground floor utility area adjacent to chlorination room",
    ],
    "winter_min_max_temperature": [
        "-2 to 42 deg C", "5 to 38 deg C", "0 to 45 deg C", "-5 to 40 deg C",
        "8 to 36 deg C", "2 to 44 deg C",
    ],
    "application": [
        "Water treatment disinfection", "Effluent treatment", "Pulp bleaching",
        "Product chlorination", "Bleaching process", "Sulfonation process",
        "Sugar refining - decolorization", "Wastewater dechlorination",
        "Textile bleaching", "Food-grade preservation",
    ],
    "orientation": ["Vertical (Bayonet type)", "Horizontal"],
    "design_capacity_kg_hr": [
        "25 kg/hr", "30 kg/hr", "35 kg/hr", "40 kg/hr", "50 kg/hr", "65 kg/hr",
        "75 kg/hr", "100 kg/hr", "120 kg/hr", "150 kg/hr", "180 kg/hr", "200 kg/hr",
        "250 kg/hr", "300 kg/hr", "350 kg/hr", "400 kg/hr", "450 kg/hr", "500 kg/hr",
        "600 kg/hr", "700 kg/hr", "750 kg/hr", "900 kg/hr",
    ],
    "heating_media": [
        "Steam", "Electric heater", "Steam (existing header available near battery limit)",
        "Electric heating jacket, high-capacity rated",
    ],
    "heating_fluid": ["Steam", "Hot water", "Thermic oil"],
    "no_of_tonners_connected": [
        "1", "2", "4", "6", "0 - fed directly from bulk storage tank",
    ],
    "tonner_manifold_arrangement_required": ["Required", "Not required"],
    "tonner_manifold_working_standby_count": [
        "2 working + 1 standby", "1 working + 1 standby", "3 working + 2 standby",
        "Not applicable",
    ],
    "instrument_specification": [
        "FLP (Zone 1 classified area)", "Non-FLP", "FLP (Zone 2 classified area)",
    ],
    "electrical_panel_specification": ["FLP", "Non-FLP"],
    "area_classification": ["Hazardous - Zone 1", "Hazardous - Zone 2", "Non-Hazardous"],
    "scope_bare_or_complete_skid": [
        "Bare vaporizer only", "Complete skid system with all instruments & controls",
    ],
    "site_layout_provided": [
        "Yes, GA drawing attached", "No, will share later", "Yes, attached as Annexure A",
    ],
    "tentative_finalization_month_year": [
        "March 2026", "June 2026", "August 2026", "December 2026", "Q1 2027", "Q3 2026",
    ],
}

GAS_SPECIFIC_POOLS = {
    "SO2": {
        "source_of_so2": [
            "SO2 tonners, 900 kg capacity", "SO2 cylinders, 68 kg capacity",
            "Bulk SO2 storage tank, 10 MT capacity", "Tonner yard, 2 tonners in use",
        ],
        "inlet_liquid_so2_temperature_c": [
            "Ambient", "Room temperature (approx 25-30 deg C)", "28 deg C", "32 deg C",
            "Normal, no cooling required",
        ],
        "inlet_liquid_so2_pressure_barg": [
            "6.5 bar (g)", "6.8 bar (g)", "7.2 bar (g)", "As per source tonner pressure",
            "Regulated to 6 bar (g)",
        ],
        "so2_pressure_at_vaporizer_outlet_barg": [
            "2.5 bar (g)", "3 bar (g)", "3.5 bar (g)", "3.8 bar (g)", "4 bar (g)",
        ],
        "so2_temperature_at_vaporizer_outlet_c": [
            "55 deg C", "60 deg C", "65 deg C", "70 deg C",
        ],
    },
    "CL2": {
        "source_of_cl2": [
            "90 kg cylinders", "900 kg tonners", "Bulk chlorine storage tank, 18 MT capacity",
            "Tonner yard, 2 tonners in use",
        ],
        "inlet_liquid_cl2_temperature_c": [
            "Ambient", "Room temperature (approx 25-30 deg C)", "28 deg C", "30 deg C",
            "Normal, no cooling required",
        ],
        "inlet_liquid_cl2_pressure_barg": [
            "6.5 bar (g)", "6.8 bar (g)", "7 bar (g)", "As per source tonner pressure",
            "Regulated to 6.5 bar (g)",
        ],
        "cl2_pressure_at_vaporizer_outlet_barg": [
            "3 bar (g)", "3.5 bar (g)", "3.8 bar (g)", "4 bar (g)",
        ],
        "cl2_temperature_at_vaporizer_outlet_c": [
            "70 deg C", "75 deg C", "80 deg C",
        ],
    },
}

FIELD_DEFS = {"SO2": SO2_TEMPLATE_FIELDS, "CL2": CL2_TEMPLATE_FIELDS}


def field_pool_for(equipment_type: str) -> dict:
    pool = dict(COMMON_POOLS)
    pool.update(GAS_SPECIFIC_POOLS[equipment_type])
    pool["description"] = [
        f"{equipment_type} Vaporizer", f"{equipment_type} Vaporizer System",
        f"{equipment_type} Vaporizer Skid", f"{equipment_type} Gas Vaporizer Unit",
    ]
    return pool


def random_spec(equipment_type: str) -> dict:
    """Returns a random SUBSET of field values — real enquiries rarely state
    every field, so completeness varies 35%-75% per sample."""
    pool = field_pool_for(equipment_type)
    all_keys = [key for key, _ in FIELD_DEFS[equipment_type]]
    fraction = random.uniform(0.35, 0.75)
    n_included = max(3, round(len(all_keys) * fraction))
    included_keys = set(random.sample(all_keys, n_included))
    return {key: random.choice(pool[key]) for key in all_keys if key in included_keys}


# ---------------------------------------------------------------------------
# Prose rendering (letter-style enquiry, fields mentioned inline)
# ---------------------------------------------------------------------------


def build_prose_body(equipment_type: str, spec: dict, field_labels: dict) -> str:
    intro = (
        f"We are evaluating options for a new {equipment_type} Vaporizer system for our "
        f"facility and would like to request your technical proposal and quotation. "
        f"Key requirements as finalized on our end so far are noted below."
    )
    clauses = [f"{field_labels[key]}: {value}" for key, value in spec.items()]
    body = intro + "\n\n" + " ".join(f"{c}." for c in clauses)
    closing = (
        "\n\nAny parameters not mentioned above are still under internal review and will be "
        "shared in a follow-up. Kindly share your technical datasheet, delivery timeline, and "
        "commercial offer at the earliest."
    )
    return body + closing


# ---------------------------------------------------------------------------
# PDF / DOCX / XLSX writers
# ---------------------------------------------------------------------------


def _mc(pdf: FPDF, h: int, text: str):
    """multi_cell wrapper — fpdf2's own default leaves the cursor at the
    right margin afterward (not the left), which then breaks the *next*
    multi_cell call with "Not enough horizontal space"."""
    pdf.multi_cell(0, h, text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)


def _letterhead_lines(equipment_type: str, customer: dict) -> list:
    return [
        f"Enquiry for Supply of {equipment_type} Vaporizer System",
        f"Enquiry Ref: {customer['enquiry_reference']}",
        "",
        customer["company"],
        customer["location"],
        f"Contact: {customer['contact_person']} ({customer['contact_title']})",
        f"Email: {customer['contact_email']}  |  Phone: {customer['contact_phone']}",
        "",
    ]


def write_pdf(path: Path, equipment_type: str, customer: dict, spec: dict, field_labels: dict, style: str):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 13)
    for line in _letterhead_lines(equipment_type, customer)[:1]:
        _mc(pdf, 8, line)
    pdf.set_font("Helvetica", size=10)
    for line in _letterhead_lines(equipment_type, customer)[1:]:
        if line:
            _mc(pdf, 6, line)
        else:
            pdf.ln(6)
    pdf.ln(2)

    if style == "table":
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(90, 7, "Parameter", border=1)
        pdf.cell(90, 7, "Details", border=1)
        pdf.ln()
        pdf.set_font("Helvetica", size=9)
        for key, value in spec.items():
            pdf.cell(90, 7, field_labels[key][:55], border=1)
            pdf.cell(90, 7, str(value)[:55], border=1)
            pdf.ln()
        pdf.ln(4)
        pdf.set_font("Helvetica", size=9)
        _mc(pdf, 6, "Kindly share your technical datasheet, delivery timeline, and commercial offer.")
    else:
        pdf.set_font("Helvetica", size=10)
        _mc(pdf, 6, build_prose_body(equipment_type, spec, field_labels))

    pdf.output(str(path))


def write_docx(path: Path, equipment_type: str, customer: dict, spec: dict, field_labels: dict, style: str):
    doc = Document()
    doc.add_heading(f"Enquiry for Supply of {equipment_type} Vaporizer System", level=1)
    for line in _letterhead_lines(equipment_type, customer)[1:]:
        if line:
            doc.add_paragraph(line)

    if style == "table":
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text = "Parameter", "Details"
        for key, value in spec.items():
            row = table.add_row().cells
            row[0].text = field_labels[key]
            row[1].text = str(value)
        doc.add_paragraph("")
        doc.add_paragraph("Kindly share your technical datasheet, delivery timeline, and commercial offer.")
    else:
        doc.add_paragraph(build_prose_body(equipment_type, spec, field_labels))

    doc.save(str(path))


def write_xlsx(path: Path, equipment_type: str, customer: dict, spec: dict, field_labels: dict, style: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RFQ"

    ws["A1"] = f"Enquiry for Supply of {equipment_type} Vaporizer System"
    ws["A2"] = "Company"
    ws["B2"] = customer["company"]
    ws["A3"] = "Location"
    ws["B3"] = customer["location"]
    ws["A4"] = "Contact Person"
    ws["B4"] = f"{customer['contact_person']} ({customer['contact_title']})"
    ws["A5"] = "Email"
    ws["B5"] = customer["contact_email"]
    ws["A6"] = "Phone"
    ws["B6"] = customer["contact_phone"]
    ws["A7"] = "Enquiry Ref"
    ws["B7"] = customer["enquiry_reference"]

    row = 9
    ws.cell(row=row, column=1, value="Parameter")
    ws.cell(row=row, column=2, value="Value")
    row += 1
    for key, value in spec.items():
        ws.cell(row=row, column=1, value=field_labels[key])
        ws.cell(row=row, column=2, value=str(value))
        row += 1

    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 45
    wb.save(str(path))


WRITERS = {"pdf": write_pdf, "docx": write_docx, "xlsx": write_xlsx}


# ---------------------------------------------------------------------------
# Main generation loop
# ---------------------------------------------------------------------------


def generate_for_type(equipment_type: str):
    field_labels = dict(FIELD_DEFS[equipment_type])
    out_root = OUTPUT_DIR / equipment_type
    for fmt in WRITERS:
        (out_root / fmt).mkdir(parents=True, exist_ok=True)

    fmt_sequence = (
        ["pdf"] * FORMAT_COUNTS["pdf"] + ["docx"] * FORMAT_COUNTS["docx"] + ["xlsx"] * FORMAT_COUNTS["xlsx"]
    )
    random.shuffle(fmt_sequence)

    manifest = []
    for i, fmt in enumerate(fmt_sequence, start=1):
        customer = random_customer(i)
        spec = random_spec(equipment_type)
        style = random.choice(["table", "prose"]) if fmt != "xlsx" else "table"
        capacity = spec.get("design_capacity_kg_hr", "unknown").replace(" ", "").replace("/", "")
        filename = f"{equipment_type.lower()}_rfq_{i:03d}_{capacity}_{style}.{fmt}"
        path = out_root / fmt / filename

        # Windows Explorer/antivirus sometimes briefly locks a just-created
        # Office file (thumbnail/preview generation) — retry past that.
        for attempt in range(5):
            try:
                WRITERS[fmt](path, equipment_type, customer, spec, field_labels, style)
                break
            except PermissionError:
                if attempt == 4:
                    raise
                time.sleep(0.5)

        manifest.append({
            "project_id": filename,
            "format": fmt,
            "style": style,
            "customer": customer,
            "fields": spec,
        })

    (out_root / "manifest.json").write_text(json.dumps(manifest, indent=2))
    print(f"{equipment_type}: generated {len(fmt_sequence)} files + manifest.json under {out_root}")


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for equipment_type in ("SO2", "CL2"):
        generate_for_type(equipment_type)
    print(f"\nDone. Files written under {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
